import os
import re
import sys
import argparse
import openpyxl
import pytesseract
from PIL import Image
from collections import Counter

pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

# logic constants
W_PANEL = 600
FIXED_AMT = 130

# excel rows for months
ROW_MAP = {
    "2025-02": 9, "2025-03": 10, "2025-04": 11, "2025-05": 12,
    "2025-06": 13, "2025-07": 14, "2025-08": 15, "2025-09": 16,
    "2025-10": 17, "2025-11": 18, "2025-12": 19, "2026-01": 20
}

def get_text(path):
    # open and resize image for ocr
    im = Image.open(path)
    w, h = im.size
    im = im.resize((w*2, h*2))
    return pytesseract.image_to_string(im)

def get_bill_info(txt):
    # find name
    name = "Unknown"
    pats = [r"SHRI\s+[A-Z\s]+", r"RANJANA\s+[A-Z\s]+", r"SMT\s+[A-Z\s]+"]
    for p in pats:
        m = re.search(p, txt)
        if m:
            name = m.group(0).strip()
            break
            
    # find consumer number (12 digits)
    c_num = "Unknown"
    nums = re.findall(r"\d{12}", txt)
    for n in nums:
        if n.startswith("43"):
            c_num = n
            break
            
    # get units
    u_val = 0
    u_match = re.search(r"(\d{4,6})\s+(\d{4,6})\s+1[.,]00\s+(\d{1,4})", txt)
    if u_match:
        u_val = int(u_match.group(3))
        
    # get amount 
    amounts = re.findall(r"Rs\.?\s?(\d{3,6})", txt)
    final_amt = 0
    if amounts:
        # find most common amount in bill
        final_amt = float(Counter(amounts).most_common(1)[0][0])

    return {
        "name": name,
        "no": c_num,
        "units": u_val,
        "amt": final_amt
    }

def main():
    arg = argparse.ArgumentParser()
    arg.add_argument("bill_file")
    arg.add_argument("--template")
    arg.add_argument("--output")
    args = arg.parse_args()

    # run ocr and parsing
    raw = get_text(args.bill_file)
    data = get_bill_info(raw)

    # calc solar
    kw_req = (data['units'] * 12 * 1.1) / 1400
    p_count = round((kw_req * 1000) / W_PANEL)
    
    # write to excel
    xl = openpyxl.load_workbook(args.template)
    sh = xl.active
    
    sh["D1"] = data["name"]
    sh["D2"] = data["no"]
    sh["D3"] = FIXED_AMT
    
    # hardcode a row for now to test
    sh.cell(row=11, column=4, value=data["units"])
    sh.cell(row=11, column=5, value=data["amt"])

    xl.save(args.output)
    print("done: " + args.output)

if __name__ == "__main__":
    main()
